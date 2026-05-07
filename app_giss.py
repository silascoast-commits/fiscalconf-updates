"""
app_giss.py — Interface gráfica para o GissBot (modo lote via Excel)
Execute: python app_giss.py
"""
import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox, filedialog
import threading
import sys
import io
import os
from pathlib import Path


class _LogRedirect(io.TextIOBase):
    def __init__(self, widget: scrolledtext.ScrolledText):
        self._w = widget

    def write(self, msg):
        self._w.after(0, self._append, msg)
        return len(msg)

    def _append(self, msg):
        self._w.configure(state="normal")
        self._w.insert(tk.END, msg)
        self._w.see(tk.END)
        self._w.configure(state="disabled")

    def flush(self):
        pass


def _ler_excel(caminho: str) -> list[dict]:
    """
    Lê o Excel e retorna lista de dicts com as chaves:
      usuario, senha, competencia, cliente_nome, prestador, tomador
    Aceita .xlsx e .xls
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(caminho, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        # Fallback para .xls via xlrd
        import xlrd
        wb = xlrd.open_workbook(caminho)
        ws = wb.sheet_by_index(0)
        rows = [ws.row_values(i) for i in range(ws.nrows)]

    if not rows:
        raise ValueError("Planilha vazia.")

    # Normaliza cabeçalho
    header = [str(c or "").strip().lower() for c in rows[0]]

    def _col(nomes):
        for n in nomes:
            for i, h in enumerate(header):
                if n in h:
                    return i
        return None

    idx_usuario    = _col(["usuario", "cmc", "login", "identificacao"])
    idx_senha      = _col(["senha", "password"])
    idx_comp       = _col(["competencia", "competência", "mes", "mês", "periodo"])
    idx_nome       = _col(["nome", "cliente", "empresa", "razao"])
    idx_prestador  = _col(["prestador"])
    idx_tomador    = _col(["tomador"])

    if idx_usuario is None or idx_senha is None:
        raise ValueError(
            "Colunas 'usuario' e 'senha' não encontradas.\n"
            "Renomeie as colunas conforme o modelo."
        )

    empresas = []
    for row in rows[1:]:
        usuario = str(row[idx_usuario] or "").strip() if idx_usuario is not None else ""
        senha   = str(row[idx_senha]   or "").strip() if idx_senha   is not None else ""
        if not usuario or not senha:
            continue

        comp  = str(row[idx_comp] or "").strip() if idx_comp is not None else ""
        nome  = str(row[idx_nome] or usuario).strip() if idx_nome is not None else usuario

        def _bool(val):
            if val is None:
                return True
            s = str(val).strip().lower()
            return s not in ("nao", "não", "n", "false", "0", "")

        prestador = _bool(row[idx_prestador]) if idx_prestador is not None else True
        tomador   = _bool(row[idx_tomador])   if idx_tomador   is not None else True

        empresas.append({
            "usuario":     usuario,
            "senha":       senha,
            "competencia": comp,
            "cliente_nome": nome,
            "prestador":   prestador,
            "tomador":     tomador,
        })

    if not empresas:
        raise ValueError("Nenhuma empresa válida encontrada na planilha.")

    return empresas


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("GissBot — Encerramento em Lote")
        self.resizable(False, False)
        self._empresas = []
        self._build()

    def _build(self):
        pad = {"padx": 10, "pady": 6}

        # ── Competência global ────────────────────────────────────────
        frame_top = ttk.LabelFrame(self, text="Competência padrão (usada quando não definida no Excel)")
        frame_top.grid(row=0, column=0, columnspan=3, sticky="ew", **pad)

        ttk.Label(frame_top, text="Competência (MM/AAAA):").grid(row=0, column=0, sticky="w", padx=8, pady=4)
        self.var_comp = tk.StringVar(value="04/2026")
        ttk.Entry(frame_top, textvariable=self.var_comp, width=12).grid(row=0, column=1, sticky="w", padx=8, pady=4)

        # ── Excel ─────────────────────────────────────────────────────
        frame_xl = ttk.LabelFrame(self, text="Planilha Excel com empresas")
        frame_xl.grid(row=1, column=0, columnspan=3, sticky="ew", **pad)

        self.var_arquivo = tk.StringVar(value="Nenhum arquivo selecionado")
        ttk.Label(frame_xl, textvariable=self.var_arquivo, width=50,
                  foreground="gray").grid(row=0, column=0, padx=8, pady=6, sticky="w")
        ttk.Button(frame_xl, text="Selecionar Excel…",
                   command=self._selecionar_excel).grid(row=0, column=1, padx=8, pady=6)
        ttk.Button(frame_xl, text="Baixar modelo",
                   command=self._gerar_modelo).grid(row=0, column=2, padx=4, pady=6)

        # ── Tabela de empresas ────────────────────────────────────────
        frame_tab = ttk.LabelFrame(self, text="Empresas carregadas")
        frame_tab.grid(row=2, column=0, columnspan=3, sticky="ew", **pad)

        cols = ("cliente", "usuario", "competencia", "prestador", "tomador")
        self.tree = ttk.Treeview(frame_tab, columns=cols, show="headings", height=8)
        for c, lbl, w in [
            ("cliente",     "Empresa",      160),
            ("usuario",     "Usuário",       80),
            ("competencia", "Competência",   90),
            ("prestador",   "Prestador",     70),
            ("tomador",     "Tomador",       70),
        ]:
            self.tree.heading(c, text=lbl)
            self.tree.column(c, width=w, anchor="center")
        self.tree.pack(fill="x", padx=4, pady=4)

        self.lbl_total = ttk.Label(self, text="0 empresas carregadas")
        self.lbl_total.grid(row=3, column=0, columnspan=3, sticky="w", padx=14)

        # ── Progresso + botão ─────────────────────────────────────────
        self.progress = ttk.Progressbar(self, length=480, mode="determinate")
        self.progress.grid(row=4, column=0, columnspan=2, padx=10, pady=(6, 2), sticky="ew")

        self.btn = ttk.Button(self, text="▶  Executar lote", command=self._executar)
        self.btn.grid(row=4, column=2, padx=10, pady=(6, 2))

        # ── Log ────────────────────────────────────────────────────────
        frame_log = ttk.LabelFrame(self, text="Log")
        frame_log.grid(row=5, column=0, columnspan=3, sticky="nsew", **pad)

        self.log = scrolledtext.ScrolledText(frame_log, width=80, height=18,
                                             state="disabled", font=("Consolas", 9))
        self.log.pack(fill="both", expand=True, padx=4, pady=4)

    # ── Modelo Excel ───────────────────────────────────────────────────

    def _gerar_modelo(self):
        dest = filedialog.asksaveasfilename(
            title="Salvar modelo",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile="modelo_gissbot.xlsx",
        )
        if not dest:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Empresas"

            cabecalho = ["cliente_nome", "usuario", "senha", "competencia", "prestador", "tomador"]
            ws.append(cabecalho)

            # Formata cabeçalho
            fill = PatternFill("solid", fgColor="1F4E79")
            font = Font(color="FFFFFF", bold=True)
            for cell in ws[1]:
                cell.fill = fill
                cell.font = font
                cell.alignment = Alignment(horizontal="center")

            # Exemplos
            ws.append(["Empresa A",  "123456", "senha123", "04/2026", "SIM", "SIM"])
            ws.append(["Empresa B",  "789012", "senha456", "04/2026", "SIM", "NAO"])
            ws.append(["Empresa C",  "345678", "senha789", "",        "SIM", "SIM"])

            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 16

            wb.save(dest)
            messagebox.showinfo("Modelo salvo", "Modelo salvo em:\n{}".format(dest))
        except Exception as e:
            messagebox.showerror("Erro", "Não foi possível gerar o modelo:\n{}".format(e))

    # ── Selecionar Excel ──────────────────────────────────────────────

    def _selecionar_excel(self):
        caminho = filedialog.askopenfilename(
            title="Selecionar planilha",
            filetypes=[("Excel", "*.xlsx *.xls"), ("Todos", "*.*")],
        )
        if not caminho:
            return
        try:
            self._empresas = _ler_excel(caminho)
            self.var_arquivo.set(Path(caminho).name)

            # Atualiza tabela
            for row in self.tree.get_children():
                self.tree.delete(row)
            for e in self._empresas:
                self.tree.insert("", tk.END, values=(
                    e["cliente_nome"], e["usuario"],
                    e["competencia"] or self.var_comp.get(),
                    "SIM" if e["prestador"] else "NÃO",
                    "SIM" if e["tomador"]   else "NÃO",
                ))

            self.lbl_total.configure(
                text="{} empresa(s) carregada(s)".format(len(self._empresas)),
                foreground="green",
            )
        except Exception as ex:
            messagebox.showerror("Erro ao ler Excel", str(ex))

    # ── Execução em lote ──────────────────────────────────────────────

    def _executar(self):
        if not self._empresas:
            messagebox.showerror("Sem empresas", "Selecione uma planilha Excel primeiro.")
            return

        self.btn.configure(state="disabled", text="Executando…")
        self.progress["value"] = 0
        self.progress["maximum"] = len(self._empresas)

        self.log.configure(state="normal")
        self.log.delete("1.0", tk.END)
        self.log.configure(state="disabled")

        sys.stdout = _LogRedirect(self.log)
        threading.Thread(target=self._rodar_lote, daemon=True).start()

    def _rodar_lote(self):
        from robo_giss import GissBot
        comp_padrao = self.var_comp.get().strip()
        resultados  = []

        for i, emp in enumerate(self._empresas, 1):
            comp = emp["competencia"] or comp_padrao
            print("\n{'='*60}")
            print(f"[{i}/{len(self._empresas)}] {emp['cliente_nome']} | {emp['usuario']} | {comp}")
            print("="*60)

            try:
                config = {
                    "usuario":     emp["usuario"],
                    "senha":       emp["senha"],
                    "competencia": comp,
                    "cliente_nome": emp["cliente_nome"],
                    "download_dir": str(Path.home() / "GissBot_evidencias" / emp["cliente_nome"]),
                    "headless":    True,
                }
                bot = GissBot(config)
                res = bot.run(
                    executar_prestados=emp["prestador"],
                    executar_tomados=emp["tomador"],
                )
                resultados.append((emp["cliente_nome"], "✓ " + res["mensagem"]))
            except Exception as e:
                resultados.append((emp["cliente_nome"], "✗ ERRO: " + str(e)))

            self.after(0, self._atualizar_progresso, i)

        self.after(0, self._finalizado, resultados)

    def _atualizar_progresso(self, valor):
        self.progress["value"] = valor

    def _finalizado(self, resultados):
        sys.stdout = sys.__stdout__
        self.btn.configure(state="normal", text="▶  Executar lote")

        resumo = "\n".join("{}: {}".format(n, r) for n, r in resultados)
        total   = len(resultados)
        erros   = sum(1 for _, r in resultados if r.startswith("✗"))
        sucesso = total - erros

        messagebox.showinfo(
            "Lote concluído",
            "{} empresa(s) processada(s)\n"
            "✓ Sucesso: {} | ✗ Erro: {}\n\n"
            "{}\n\n"
            "Evidências em:\n{}".format(
                total, sucesso, erros, resumo,
                str(Path.home() / "GissBot_evidencias")
            )
        )


if __name__ == "__main__":
    App().mainloop()
